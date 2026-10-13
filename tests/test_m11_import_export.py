# -*- coding: utf-8 -*-
"""M11 导入导出测试（L5）。
覆盖：项目/资产模板下载与权限、导入执行（成功/部分失败/编号重复）、导出表头与行数=列表口径、
导出筛选一致、导出按权限过滤、导出往返一致性（抽样比对库内）、空数据集导出。
"""
import os
import sys
import datetime
from io import BytesIO

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _common as C  # noqa: E402
import _oracle as O  # noqa: E402
import openpyxl  # noqa: E402


def main():
    C.setup_stdout()
    C.ensure_dirs()
    C.prepare_runtime()
    C.prepare_base()
    db = C.work_db("m11_ie.db")
    prod_before = C.prod_fingerprint()
    r = C.R("M11", "M11 导入导出测试报告（L5）")
    srv = C.Server(db)
    srv.start()
    B = srv.base
    L = C.RUNTIME_DB
    stamp = datetime.datetime.now().strftime("%H%M%S")
    try:
        A, ia = C.login("admin", base=B)
        assert A, "admin 登录失败"
        p_un = C.query(L, "SELECT 用户名 FROM users WHERE 角色='普通用户' ORDER BY 用户名 LIMIT 1")[0]["用户名"]
        P, ip = C.login(p_un, base=B)
        q_un = C.query(L, "SELECT 用户名 FROM users WHERE 角色='查询用户' ORDER BY 用户名 LIMIT 1")[0]["用户名"]
        Q, iq = C.login(q_un, base=B)

        def xlsx(resp):
            assert resp.status_code == 200, "HTTP %s %s" % (resp.status_code, resp.text[:120])
            return openpyxl.load_workbook(BytesIO(resp.content)).active

        def import_file(rows, headers=None):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "合同台账"
            ws.append(headers or ["编号", "项目分类", "项目名称", "采购方式", "合同金额（万元）", "资金来源",
                                  "合同主要内容", "经办人", "所属部门", "是否结算"])
            for x in rows:
                ws.append(x)
            bio = BytesIO()
            wb.save(bio)
            return bio

        def cnt():
            return len(C.query(L, "SELECT id FROM project_contracts"))

        def t01():
            ws = xlsx(A.get(B + "/api/contracts/template", timeout=60))
            hdr = [str(c.value or "") for c in ws[1]]
            need = ["编号", "项目分类", "项目名称", "采购方式", "合同金额（万元）", "资金来源", "经办人",
                    "所属部门", "提醒提前天数"]
            miss = [x for x in need if x not in hdr]
            assert not miss, "模板缺少列：%s" % miss
            return "模板 %d 列、示例行 %d 行；必填列齐全" % (len(hdr), max(0, ws.max_row - 1))

        r.case("TC-M11-01", "项目导入模板下载（管理员）", "200 + xlsx + 表头齐全", t01)

        def t02():
            hits = []
            for name, sess in (("普通用户", P), ("查询用户", Q)):
                for path in ("/api/contracts/template", "/api/asset_contracts/template"):
                    rr = sess.get(B + path, timeout=30)
                    hits.append("%s/%s=%s" % (name, path.rsplit("/", 1)[-1], rr.status_code))
                    assert rr.status_code == 403, "%s 访问 %s 返回 %s（期望 403）" % (name, path, rr.status_code)
            return "非管理员访问模板均 403（%s）" % ",".join(hits)

        r.case("TC-M11-02", "模板下载权限", "非管理员 403", t02)

        def t03():
            ws = xlsx(A.get(B + "/api/asset_contracts/template", timeout=60))
            hdr = [str(c.value or "") for c in ws[1]]
            need = ["位置", "资产名称", "承租人", "管理部门", "年租金（元）", "租期开始日", "下次出租时间"]
            miss = [x for x in need if x not in hdr]
            assert not miss, "资产模板缺少列：%s" % miss
            return "资产模板 %d 列；关键列齐全" % len(hdr)

        r.case("TC-M11-03", "资产导入模板下载", "200 + 表头含资产必填列", t03)

        def t04():
            no = "IMPEXEC" + stamp[:2]
            f = import_file([[no, "服务类", "导入执行-" + stamp, "自行采购", 3.5, "其他", "x",
                              ia["name"], ia["department"], "否"]])
            before = cnt()
            resp = A.post(B + "/api/contracts/import", files={"file": ("imp.xlsx", f.getvalue())}, timeout=60)
            assert resp.status_code == 200, "导入 HTTP %s %s" % (resp.status_code, resp.text[:160])
            after = cnt()
            assert after == before + 1, "导入未新增（%d→%d）" % (before, after)
            row = C.query(L, "SELECT * FROM project_contracts WHERE 编号=?", (no,))
            assert row and abs(O.num(row[0]["合同金额（万元）"]) - 3.5) < 0.001, "导入数据不符"
            return "导入 1 行成功：合同数 %d→%d，编号 %s 金额 3.5 落库正确" % (before, after, no)

        r.case("TC-M11-04", "导入执行（有效行）", "200 + 库内新增", t04)

        def t05():
            dup = C.query(L, "SELECT 编号 n FROM project_contracts ORDER BY id LIMIT 1")[0]["n"]
            newno = "IMPGOOD" + stamp[:2]
            f = import_file([[newno, "服务类", "导入部分失败-" + stamp, "自行采购", 1, "其他", "x",
                              ia["name"], ia["department"], "否"],
                             [dup, "服务类", "重复编号行", "自行采购", 1, "其他", "x",
                              ia["name"], ia["department"], "否"]])
            before = cnt()
            resp = A.post(B + "/api/contracts/import", files={"file": ("imp2.xlsx", f.getvalue())}, timeout=60)
            assert resp.status_code == 200, "导入 HTTP %s" % resp.status_code
            after = cnt()
            assert after == before + 1, "有效行未入库或错误行被写入（%d→%d）" % (before, after)
            assert C.query(L, "SELECT id FROM project_contracts WHERE 编号=?", (newno,)), "有效行缺失"
            assert dup in resp.text or "已存在" in resp.text, "未回显错误行原因：%s" % resp.text[:200]
            return "1 有效行入库、1 重复编号行被跳过并回显原因（%s）" % dup

        r.case("TC-M11-05", "导入部分失败（错误行跳过并回显）", "仅有效行入库", t05, priority="P1")

        def t06():
            dup = C.query(L, "SELECT 编号 n FROM project_contracts ORDER BY id LIMIT 1")[0]["n"]
            f = import_file([[dup, "服务类", "重复编号-单独", "自行采购", 1, "其他", "x",
                              ia["name"], ia["department"], "否"]])
            before = cnt()
            resp = A.post(B + "/api/contracts/import", files={"file": ("imp3.xlsx", f.getvalue())}, timeout=60)
            after = cnt()
            assert after == before, "重复编号行被写入（%d→%d）" % (before, after)
            assert resp.status_code == 200, "重复编号导入 HTTP %s" % resp.status_code
            return "重复编号行未入库（库内仍 %d 条），响应回显原因" % after

        r.case("TC-M11-06", "导入编号重复拒绝", "不写入 + 回显", t06, priority="P1")

        def t07():
            ws = xlsx(A.get(B + "/api/contracts/export", timeout=180))
            hdr = [str(c.value or "") for c in ws[1]]
            total = A.get(B + "/api/contracts", params={"page": 1, "page_size": 1}, timeout=60).json()["total"]
            n = ws.max_row - 1
            assert n == total, "导出行数 %d ≠ 列表 total %d" % (n, total)
            for k in ("编号", "项目名称", "合同金额（万元）", "所属部门"):
                assert k in hdr, "导出缺少列 %s" % k
            return "导出 %d 行 = 列表 total %d；表头 %d 列含关键列" % (n, total, len(hdr))

        r.case("TC-M11-07", "项目导出表头与行数", "行数=列表 total", t07)

        def t08():
            cat = "服务类"
            ws = xlsx(A.get(B + "/api/contracts/export", params={"category": cat}, timeout=180))
            exp = A.get(B + "/api/contracts", params={"category": cat, "page": 1, "page_size": 1}, timeout=60).json()["total"]
            assert ws.max_row - 1 == exp, "筛选导出 %d 行 ≠ 列表 %d" % (ws.max_row - 1, exp)
            return "category=%s：导出 %d 行 = 列表 total" % (cat, exp)

        r.case("TC-M11-08", "导出筛选与列表一致", "行数一致", t08)

        def t09():
            ws = xlsx(P.get(B + "/api/contracts/export", timeout=180))
            exp = P.get(B + "/api/contracts", params={"page": 1, "page_size": 1}, timeout=60).json()["total"]
            admin = A.get(B + "/api/contracts", params={"page": 1, "page_size": 1}, timeout=60).json()["total"]
            assert ws.max_row - 1 == exp, "普通用户导出 %d 行 ≠ 其可见 total %d" % (ws.max_row - 1, exp)
            assert exp < admin, "普通用户可见数未收敛"
            return "普通用户导出 %d 行 = 其可见范围（管理员 %d 行）" % (exp, admin)

        r.case("TC-M11-09", "导出按数据权限过滤", "仅导出可见范围", t09)

        def t10():
            ws = xlsx(A.get(B + "/api/contracts/export", timeout=180))
            hdr = [str(c.value or "") for c in ws[1]]
            idx = {h: i for i, h in enumerate(hdr)}
            db_rows = {str(x["编号"]): x for x in C.query(L, "SELECT * FROM project_contracts")}
            checked = bad = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                no = str(row[idx["编号"]] or "").strip()
                if no not in db_rows:
                    continue
                d = db_rows[no]
                amt = row[idx["合同金额（万元）"]]
                exp_amt = O.num(d["合同金额（万元）"])
                if exp_amt is not None and (amt is None or abs(float(amt) - exp_amt) > 0.01):
                    bad += 1
                if str(row[idx["所属部门"]] or "").strip() != str(d["所属部门"] or "").strip():
                    bad += 1
                if str(row[idx["项目名称"]] or "").strip() != str(d["项目名称"] or "").strip():
                    bad += 1
                checked += 1
                if checked >= 30:
                    break
            assert checked >= 20, "抽样行数不足（%d）" % checked
            assert bad == 0, "往返比对不一致 %d 处" % bad
            return "抽样 %d 条：编号↔项目名称/合同金额/所属部门 全部与库内一致" % checked

        r.case("TC-M11-10", "项目导出往返一致性", "抽样字段一致", t10, priority="P1")

        def t11():
            resp = A.get(B + "/api/asset_contracts/export", timeout=120)
            assert resp.status_code == 200, "资产导出 HTTP %s %s" % (resp.status_code, resp.text[:120])
            ws = openpyxl.load_workbook(BytesIO(resp.content)).active
            hdr = [str(c.value or "") for c in ws[1]]
            total = A.get(B + "/api/asset_contracts", params={"page": 1, "page_size": 1}, timeout=60).json()["total"]
            assert ws.max_row - 1 == total, "资产导出行数 %d ≠ 列表 %d" % (ws.max_row - 1, total)
            return "资产导出 %d 行 = 列表 total %d；表头 %d 列" % (ws.max_row - 1, total, len(hdr))

        r.case("TC-M11-11", "资产导出与列表一致", "行数一致", t11, priority="P1")

        def t12():
            unit = "测试单位" + stamp
            deptn = "空数据部门" + stamp
            d1 = A.post(B + "/api/departments", json={"unit": unit, "dept": deptn}, timeout=30)
            assert d1.status_code == 201, "建部门失败 %s %s" % (d1.status_code, d1.text[:120])
            un = "zz_empty_%s" % stamp
            u1 = A.post(B + "/api/users", json={"username": un, "name": "空数据用户",
                                                "department": "%s-%s" % (unit, deptn), "role": "普通用户"}, timeout=30)
            assert u1.status_code == 201, "建用户失败 %s %s" % (u1.status_code, u1.text[:120])
            S, info = C.login(un, base=B)
            assert S, "空数据用户登录失败"
            ws = xlsx(S.get(B + "/api/contracts/export", timeout=120))
            total = S.get(B + "/api/contracts", params={"page": 1, "page_size": 1}, timeout=30).json()["total"]
            assert total == 0, "空部门用户可见 %d 条" % total
            assert ws.max_row == 1, "空数据集导出应为仅表头，实际 %d 行" % ws.max_row
            return "无可见数据的账号：列表 total=0，导出 200 且仅表头（1 行）"

        r.case("TC-M11-12", "空数据集导出", "200 + 仅表头", t12, priority="P2")

        srv.sync_db()
    finally:
        srv.stop()

    prod_after = C.prod_fingerprint()
    meta = {
        "服务地址": B,
        "工作库": "test_data/m11_ie.db",
        "生产库未改动": "是" if prod_before == prod_after else "否（指纹变化！）",
    }
    r.report("M11_导入导出测试.md", meta=meta)
    n, p, f, e = r.stat()
    print("\n[SUMMARY] M11 共 %d 例：通过 %d，失败 %d，错误 %d" % (n, p, f, e))
    sys.stdout.flush()
    os._exit(0 if (f == 0 and e == 0) else 1)


if __name__ == "__main__":
    main()
